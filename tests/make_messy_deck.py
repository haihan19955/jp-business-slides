"""テスト用: よくある「崩れた日本語資料」を python-pptx で作る（このスキルの jpdeck.js は使わない）。"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt

out = sys.argv[1] if len(sys.argv) > 1 else "messy.pptx"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
s.shapes.title.text = "受注業務の現状について。"
tb = s.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(3.4), Inches(1.0))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run()
r.text = "受注担当者が毎日ＦＡＸで届く注文書を基幹システムへ手入力しており、繁忙期には残業が月40時間を超えている。"
r.font.size = Pt(14)
r.font.name = "Calibri"
q = tf.add_paragraph()
r2 = q.add_run()
r2.text = "対応方針は来月までに決定します。"
r2.font.size = Pt(10)

rows, cols = 3, 2
tbl = s.shapes.add_table(rows, cols, Inches(4.5), Inches(1.6), Inches(4.0), Inches(1.5)).table
cells = [["項目", "内容"], ["年間削減額", "約860万円の人件費削減が見込めると試算している"], ["導入期間", "3ヶ月程度を想定しています"]]
for i in range(rows):
    for j in range(cols):
        tbl.cell(i, j).text = cells[i][j]
prs.save(out)
print("wrote", out)
