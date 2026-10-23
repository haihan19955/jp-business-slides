#!/usr/bin/env python3
"""jp_fix_wrap.py — PowerPoint の「泣き別れ」（単語の途中での改行）と禁則違反を自動で直す

  python jp_fix_wrap.py deck.pptx -o deck_fixed.pptx              # Windows + PowerPoint（実機で測る）
  python jp_fix_wrap.py deck.pptx -o deck_fixed.pptx --simulate   # Mac / Linux（文字幅から再現）

1. 日本語の段落のテキスト言語を ja-JP にする（en-US のままだと PowerPoint が禁則処理をせず、
   「、」や「）」が行頭に来る。pptxgenjs や python-pptx で作った資料は en-US のことが多い）
2. 改行位置を取得 → 文節（BudouX）の途中で切れている箇所の直前に行内改行
   （Shift+Enter と同じ <a:br/>）を入れる → もう一度測る、を繰り返す
既存の資料（このスキルで作っていない .pptx）にも使える。元ファイルは上書きしない。

必要: python-pptx、budoux。実機で測るなら Windows + PowerPoint、--simulate なら fonttools
"""
import argparse
import copy
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from jp_lint import JP_RE, find_bad_breaks  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def export_layout(pptx_path, out_json):
    cmd = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
           os.path.join(HERE, "export_layout.ps1"), "-Pptx", pptx_path, "-Out", out_json]
    subprocess.run(cmd, check=True, capture_output=True, cwd=os.path.dirname(out_json) or ".")
    with open(out_json, encoding="utf-8-sig") as fh:
        return json.load(fh)


def set_lang_ja(pptx_path):
    """日本語を含む段落の run の言語を ja-JP にする。直した run の数を返す。

    lang が en-US のときだけでなく、未指定のときも PowerPoint は日本語の禁則処理をしない
    （tests/messy.pptx で確認）ので、未指定の run にも ja-JP を書く。
    """
    prs = Presentation(pptx_path)
    changed = 0
    for slide in prs.slides:
        for p_elm in slide._element.iter(A + "p"):
            text = "".join(t.text or "" for t in p_elm.iter(A + "t"))
            if not JP_RE.search(text):
                continue
            for r in p_elm.findall(A + "r"):
                rpr = r.find(A + "rPr")
                if rpr is None:
                    rpr = r.makeelement(A + "rPr", {})
                    r.insert(0, rpr)
                lang = rpr.get("lang")
                if not lang or not lang.lower().startswith("ja"):
                    rpr.set("lang", "ja-JP")
                    rpr.set("altLang", "en-US")
                    changed += 1
    prs.save(pptx_path)
    return changed


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def find_text_frame(slide, shape_id, where):
    for sh in iter_shapes(slide.shapes):
        if sh.shape_id != shape_id:
            continue
        m = re.search(r"\[r(\d+)c(\d+)\]$", where)
        if m and getattr(sh, "has_table", False) and sh.has_table:
            return sh.table.cell(int(m.group(1)) - 1, int(m.group(2)) - 1).text_frame
        if sh.has_text_frame:
            return sh.text_frame
    return None


def insert_break(p_elm, offset):
    """段落内のテキスト位置 offset の直前に <a:br/> を入れる（run は必要なら分割）。"""
    pos = 0
    for child in list(p_elm):
        tag = child.tag
        if tag not in (A + "r", A + "fld"):
            continue
        t = child.find(A + "t")
        text = t.text or "" if t is not None else ""
        if pos + len(text) <= offset and not (pos == offset):
            pos += len(text)
            continue
        br = p_elm.makeelement(A + "br", {})
        rpr = child.find(A + "rPr")
        if rpr is not None:
            br.append(copy.deepcopy(rpr))
        k = offset - pos
        if k <= 0 or tag == A + "fld":
            child.addprevious(br)
            return True
        # run を2つに分けて間に改行を入れる
        tail = copy.deepcopy(child)
        t.text = text[:k]
        tail.find(A + "t").text = text[k:]
        child.addnext(br)
        br.addnext(tail)
        return True
    return False


def main():
    ap = argparse.ArgumentParser(description="PowerPoint の泣き別れを自動修正する")
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--rounds", type=int, default=5)
    ap.add_argument("--simulate", action="store_true",
                    help="PowerPoint を使わず、フォントの文字幅から改行位置を再現する（Mac / Linux 向け）")
    a = ap.parse_args()

    src, out = os.path.abspath(a.pptx), os.path.abspath(a.out)
    if src == out:
        sys.exit("出力先は元ファイルと別にしてください")
    shutil.copyfile(src, out)
    work = tempfile.mkdtemp(prefix="jpwrap_")
    total, unfixable = 0, []

    n_lang = set_lang_ja(out)
    if n_lang:
        print(f"  テキストの言語を ja-JP に変更（{n_lang} か所）→ PowerPoint の禁則処理が効くようになる")

    if a.simulate:
        from jp_wrap_sim import simulate_deck
        get_layout = lambda rnd: simulate_deck(out)[0]  # noqa: E731
        print("  改行位置はシミュレーション（PowerPoint 実機との一致率は約99%）")
    else:
        get_layout = lambda rnd: export_layout(out, os.path.join(work, f"layout{rnd}.json"))  # noqa: E731

    for rnd in range(1, a.rounds + 1):
        layout = get_layout(rnd)
        fixes, unfixable = [], []
        for f in layout:
            if "Slide Number" in f["shape"]:
                continue
            for p in f["paragraphs"]:
                bad = find_bad_breaks(p)
                ok = [b for b in bad if b["insert_at"] is not None]
                unfixable += [(f["slide"], f["shape"], b) for b in bad if b["insert_at"] is None]
                if ok:  # 1段落につき最初の1か所だけ直す（直すと後ろの改行位置が変わるため）
                    fixes.append((f["slide"], f["id"], f["shape"], p["index"], ok[0]))
        if not fixes:
            break
        prs = Presentation(out)
        done = 0
        for slide_no, shape_id, where, pidx, b in fixes:
            tf = find_text_frame(prs.slides[slide_no - 1], shape_id, where)
            if tf is None or pidx >= len(tf.paragraphs):
                continue
            if insert_break(tf.paragraphs[pidx]._p, b["insert_at"]):
                done += 1
                print(f"  round {rnd}: スライド{slide_no} [{where}] 「{b['phrase']}」の前で改行")
        prs.save(out)
        total += done
        if done == 0:
            break

    print(f"\n修正 {total} か所 → {out}")
    if unfixable:
        print(f"自動では直せない箇所 {len(unfixable)} 件（文節が1行より長い）:")
        for sl, where, b in unfixable:
            print(f"  スライド{sl} [{where}] 「…{b['prev'][-8:]}｜{b['next'][:8]}…」→ 言い回しか枠幅を変える")
    shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
