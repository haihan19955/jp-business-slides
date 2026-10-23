#!/usr/bin/env python3
"""jp_wrap_sim.py — PowerPoint の日本語の折り返しを、フォントの文字幅から再現する

PowerPoint がない環境（Mac / Linux）で export_layout.ps1 の代わりに使う。
出力は export_layout.ps1 と同じ形式の JSON なので、jp_lint.py --layout にそのまま渡せる。

  python jp_wrap_sim.py deck.pptx -o layout.json

再現の精度は tests/ の検証コーパス（600 ボックス）で PowerPoint 実機と比べて測っている（README 参照）。
フォントファイルが見つからない書体は、近い書体で代用し、その旨を出力する。

依存: python-pptx, fonttools
"""
import argparse
import glob
import json
import os
import re
import sys
import unicodedata
from functools import lru_cache

from fontTools.ttLib import TTCollection, TTFont
from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from jp_lint import NO_END, NO_START  # noqa: E402

EMU_PER_IN = 914400
DEFAULT_INSET_LR = 91440   # 0.1in
DEFAULT_INSET_TB = 45720   # 0.05in

HANG = set("、。，．")  # 句読点のぶら下げ（全角のみ。半角の「.」は 8.4% のような数字の一部なのでぶら下げない）

FONT_DIRS = [
    os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"),
    os.path.expanduser(r"~\AppData\Local\Microsoft\Windows\Fonts"),
    "/System/Library/Fonts", "/System/Library/Fonts/Supplemental", "/Library/Fonts",
    os.path.expanduser("~/Library/Fonts"), "/usr/share/fonts", "/usr/local/share/fonts", os.path.expanduser("~/.fonts"),
]
# 見つからないときの代用（Mac で Meiryo がない、など）
FALLBACKS = {
    "meiryo": ["yu gothic", "hiragino sans", "hiragino kaku gothic pron", "noto sans jp", "noto sans cjk jp"],
    "meiryo ui": ["meiryo", "yu gothic ui", "hiragino sans", "noto sans jp"],
    "yu gothic": ["yugothic", "hiragino sans", "noto sans jp", "meiryo"],
    "biz udpgothic": ["biz udpゴシック", "hiragino sans", "noto sans jp", "meiryo"],
    "ms pgothic": ["ｍｓ ｐゴシック", "hiragino sans", "noto sans jp"],
}


def _names(font):
    out = set()
    for rec in font["name"].names:
        if rec.nameID in (1, 4, 16):
            try:
                out.add(rec.toUnicode().strip().lower())
            except Exception:
                pass
    return out


def _weight(font):
    try:
        return font["OS/2"].usWeightClass
    except Exception:
        return 400


@lru_cache(maxsize=1)
def font_index():
    """{(family_lower, bold): (path, index)} を作る。

    同じ名前の書体が複数ある（游ゴシックの Light / Regular / Medium / Bold など）ときは、
    PowerPoint が使うウェイトに近いもの（標準なら 400、太字なら 700）を選ぶ。
    """
    found = {}
    for d in FONT_DIRS:
        if not os.path.isdir(d):
            continue
        for path in glob.glob(os.path.join(d, "**", "*"), recursive=True):
            ext = path.lower().rsplit(".", 1)[-1]
            if ext not in ("ttf", "otf", "ttc", "otc"):
                continue
            try:
                faces = TTCollection(path, lazy=True).fonts if ext in ("ttc", "otc") else [TTFont(path, lazy=True)]
            except Exception:
                continue
            for i, f in enumerate(faces):
                try:
                    names, wt = _names(f), _weight(f)
                except Exception:
                    continue
                for n in names:
                    found.setdefault(n, []).append((wt, path, i))
    idx = {}
    for n, faces in found.items():
        for bold, target in ((False, 400), (True, 700)):
            wt, path, i = min(faces, key=lambda t: (abs(t[0] - target), t[0]))
            if bold and wt < 600:
                continue  # 太字の書体がない
            idx[(n, bold)] = (path, i)
    return idx


@lru_cache(maxsize=64)
def load_metrics(family, bold):
    """(advance(ch)->em, line_height_em, resolved_name)"""
    idx = font_index()
    fam = family.strip().lower()
    cands = [fam] + FALLBACKS.get(fam, []) + ["meiryo", "yu gothic", "hiragino sans", "noto sans jp", "noto sans cjk jp", "arial"]
    for c in cands:
        for b in ((True, False) if bold else (False,)):
            key = (c, b) if (c, b) in idx else None
            if key is None:
                continue
            path, i = idx[key]
            f = TTCollection(path).fonts[i] if path.lower().endswith(("ttc", "otc")) else TTFont(path)
            upm = f["head"].unitsPerEm
            cmap = f.getBestCmap()
            hmtx = f["hmtx"].metrics
            os2 = f["OS/2"]
            asc = os2.usWinAscent
            desc = os2.usWinDescent
            cache = {}
            # 文字幅は書体の hmtx をそのまま使う。プロポーショナル書体（BIZ UDPゴシックなど）のかなも
            # PowerPoint は hmtx の幅で並べる（全角幅に置き換えると一致率が 97% → 36% に落ちる）。

            def adv(ch, cmap=cmap, hmtx=hmtx, upm=upm, cache=cache):
                if ch not in cache:
                    g = cmap.get(ord(ch))
                    if g is None:
                        cache[ch] = 1.0 if unicodedata.east_asian_width(ch) in "WF" else 0.55
                    else:
                        cache[ch] = hmtx[g][0] / upm
                return cache[ch]
            return adv, (asc + desc) / upm, c
    raise RuntimeError("日本語フォントが見つかりません")


def is_word_char(ch):
    # 欧文の単語として途中で切らない文字（英数字と一部の記号）
    return ch.isascii() and (ch.isalnum() or ch in "._%+#&@'’/:-,=") or ch in "０１２３４５６７８９"


def break_allowed(prev, cur, kinsoku=True):
    """prev と cur の間で改行してよいか（cur が次の行頭）。

    kinsoku: 禁則処理をするか。PowerPoint はテキストの言語（rPr の lang）が ja-JP のときだけ
    日本語の禁則処理をする。pptxgenjs などが書く en-US のままだと「、」や「）」が行頭に来る。
    """
    if kinsoku and (cur in NO_START or prev in NO_END):
        return False
    if prev == " ":
        return True
    if is_word_char(prev) and is_word_char(cur):
        return prev == "-" and cur.isalpha()   # ハイフンの後ろは切れる
    return True


def wrap(text, width_pt, size_pt, adv_of, kinsoku=True, hang=True):
    """text を幅 width_pt に収まるように折り返した行のリストを返す。

    adv_of(ch) は em 単位の幅。size_pt は全体の文字サイズ（float）か、文字ごとのサイズのリスト
    （「120」は54pt、「時間/月」は20pt のように run ごとにサイズが違う段落用）。
    """
    sizes = size_pt if isinstance(size_pt, list) else [size_pt] * len(text)
    width = lambda j: adv_of(text[j]) * sizes[j]  # noqa: E731
    lines, start, x = [], 0, 0.0
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        w = width(i)
        if x + w > width_pt + 0.01 and i > start:
            if kinsoku and hang and ch in HANG:  # 句読点はぶら下げて、この行に入れる
                lines.append(text[start:i + 1])
                start, x, i = i + 1, 0.0, i + 1
                continue
            # 直前の改行可能位置まで戻る
            b = i
            while b > start and not break_allowed(text[b - 1], text[b], kinsoku):
                b -= 1
            if b == start:
                b = i  # どこでも切れない（長い英単語など）→ 強制的に切る
            lines.append(text[start:b].rstrip(" "))
            start = b
            while start < n and text[start] == " ":
                start += 1
            x = sum(width(j) for j in range(start, i))
            i = max(i, start)
            continue
        x += w
        i += 1
    if start < n:
        lines.append(text[start:])
    return lines


def _attr(el, name, default):
    v = el.get(name) if el is not None else None
    return int(v) if v is not None else default


def simulate_deck(path):
    prs = Presentation(path)
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    out, notes = [], set()

    def frame(slide_no, shape_id, name, kind, tf, box_w, box_h):
        body = tf._txBody.find(A + "bodyPr")
        l, r = _attr(body, "lIns", DEFAULT_INSET_LR), _attr(body, "rIns", DEFAULT_INSET_LR)
        t, b = _attr(body, "tIns", DEFAULT_INSET_TB), _attr(body, "bIns", DEFAULT_INSET_TB)
        wrap_none = body is not None and body.get("wrap") == "none"
        vertical = body is not None and body.get("vert") not in (None, "horz")  # 縦書きは折り返しを見ない
        if not tf.text.strip():
            return  # 文字のない図形（export_layout.ps1 も出力しない）
        paras, total_h = [], 0.0
        for pi, p in enumerate(tf.paragraphs):
            runs = [(r_.text, r_) for r_ in p.runs]
            text = "".join(t_ for t_, _ in runs)
            ppr = p._p.find(A + "pPr")
            mar_l = _attr(ppr, "marL", 0) + max(0, _attr(ppr, "indent", 0))
            avail_pt = (box_w - l - r - mar_l) / 12700
            # 書体・太字・言語は段落内の最初の run で代表させる（混在時の誤差は許容）。サイズは run ごとに見る
            rr = runs[0][1] if runs else None
            run_size = lambda r_: (r_.font.size.pt if r_.font.size else None) or 18.0  # noqa: E731
            size = run_size(rr) if rr is not None else 18.0
            bold = bool(rr is not None and rr.font.bold)
            fam = None
            lang = None
            if rr is not None:
                rpr = rr._r.find(A + "rPr")
                if rpr is not None:
                    lang = rpr.get("lang")
                    for tag in ("ea", "latin"):
                        e = rpr.find(A + tag)
                        if e is not None and e.get("typeface") and not e.get("typeface").startswith("+"):
                            fam = e.get("typeface")
                            break
            fam = fam or "Meiryo"
            adv, lh_em, resolved = load_metrics(fam, bold)
            if resolved != fam.lower():
                notes.add(f"{fam} → {resolved} で代用")
            # 行内改行（<a:br/>）ごとに分けて折り返す。文字ごとのサイズも持っておく
            segs, cur, cur_sz = [], "", []
            run_of = {id(r_._r): r_ for _, r_ in runs}
            for child in p._p:
                if child.tag == A + "br":
                    segs.append((cur, cur_sz)); cur, cur_sz = "", []
                elif child.tag in (A + "r", A + "fld"):
                    tt = child.find(A + "t")
                    s_txt = (tt.text or "") if tt is not None else ""
                    r_ = run_of.get(id(child))
                    cur += s_txt
                    cur_sz += [run_size(r_) if r_ is not None else size] * len(s_txt)
            segs.append((cur, cur_sz))
            # 禁則処理は lang が ja-JP のときだけ。未指定でも禁則なしになる（tests/messy.pptx で確認）
            kin = bool(lang) and lang.lower().startswith("ja")
            spc = 1.0
            ln = ppr.find(A + "lnSpc") if ppr is not None else None
            if ln is not None and ln.find(A + "spcPct") is not None:
                spc = int(ln.find(A + "spcPct").get("val")) / 100000
            lines = []
            for sgm, sz in segs:
                if wrap_none or vertical or not sgm:
                    seg_lines = [sgm]
                else:
                    seg_lines = wrap(sgm, avail_pt, sz, adv, kinsoku=kin)
                # 行の高さはその行で一番大きい文字で決まる
                pos = 0
                for ln_txt in seg_lines:
                    k = sgm.find(ln_txt, pos) if ln_txt else pos
                    k = pos if k < 0 else k
                    line_sz = max(sz[k:k + len(ln_txt)] or [size])
                    pos = k + len(ln_txt)
                    total_h += line_sz * lh_em * spc
                lines += seg_lines
            paras.append({"index": pi, "text": text, "lines": lines, "size": size, "font": fam, "fontEA": fam, "lang": lang})
        if vertical:
            total_h = 0.0  # 縦書きの高さは見積もらない
        out.append({
            "slide": slide_no, "id": shape_id, "shape": name, "kind": kind,
            "boxW": round(box_w / EMU_PER_IN, 3), "boxH": round(box_h / EMU_PER_IN, 3),
            "textH": round(total_h / 72, 3), "innerH": round((box_h - t - b) / EMU_PER_IN, 3),
            "autoSize": 0, "wordWrap": not wrap_none, "paragraphs": paras, "simulated": True,
        })

    def walk(slide_no, shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                walk(slide_no, sh.shapes)
            elif getattr(sh, "has_table", False) and sh.has_table:
                tbl = sh.table
                for ri, row in enumerate(tbl.rows):
                    for ci, cell in enumerate(row.cells):
                        frame(slide_no, sh.shape_id, f"{sh.name}[r{ri + 1}c{ci + 1}]", "table", cell.text_frame,
                              tbl.columns[ci].width, row.height)
            elif sh.has_text_frame and sh.width is not None:
                frame(slide_no, sh.shape_id, sh.name, "text", sh.text_frame, sh.width, sh.height)

    for i, s in enumerate(prs.slides, 1):
        walk(i, s.shapes)
    return out, sorted(notes)


def main():
    ap = argparse.ArgumentParser(description="PowerPoint の日本語の折り返しをフォントの文字幅から再現する")
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out", default="layout.json")
    a = ap.parse_args()
    layout, notes = simulate_deck(a.pptx)
    with open(a.out, "w", encoding="utf-8") as fh:
        json.dump(layout, fh, ensure_ascii=False, indent=1)
    print(f"wrote {len(layout)} text frames to {a.out}（シミュレーション）")
    for n in notes:
        print("note:", n, file=sys.stderr)


if __name__ == "__main__":
    main()
